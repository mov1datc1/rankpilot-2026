import asyncio
import json
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch, AsyncMock
from zipfile import ZipFile
import fitz
from docx import Document
from docx.oxml import OxmlElement
from utils.document_preflight import read_document, SourceError
from utils.doc_parser import DocumentParser

SOURCE = 'The team represented Synthetic Client in a tax appeal. The assessment is MXN 5000 and the appeal remains pending.'

def form(label='Confidential Matter 1', summary=SOURCE):
    return f'{label}\nE1 Name of client\nSynthetic Client\nE2 Summary of matter\n{summary}\nE3 Matter value\nMXN 5000\nE5 Lead partner\nSynthetic Partner\n'

class PreflightTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
    def tearDown(self):
        self.temp.cleanup()
    def assert_error(self, path, code):
        with self.assertRaises(SourceError) as caught:
            read_document(str(path))
        self.assertEqual(caught.exception.code, code)
        return caught.exception
    def word(self, name='source.docx'):
        path=self.root/name
        doc=Document()
        for line in form().splitlines(): doc.add_paragraph(line)
        doc.save(path)
        return path
    def test_real_content_overrides_wrong_extension(self):
        path=self.word('renamed.doc')
        text, report=read_document(str(path))
        self.assertIn(SOURCE,text)
        self.assertEqual(report['detected_format'],'docx')
        self.assertTrue(report['warnings'])
    def test_xml_tables_controls_paragraphs_and_breaks_are_not_duplicated(self):
        path=self.root/'table.docx';doc=Document()
        table=doc.add_table(rows=2, cols=2)
        table.cell(0,0).text='E2 Summary of matter'
        cell=table.cell(0,1);cell.text='First paragraph.';cell.add_paragraph('Second paragraph.')
        run=cell.add_paragraph().add_run('Before');run.add_break();run.add_text('After')
        control=OxmlElement('w:sdt');content=OxmlElement('w:sdtContent');control.append(content)
        para=OxmlElement('w:p');run=OxmlElement('w:r');text=OxmlElement('w:t');text.text='Controlled text only once';run.append(text);para.append(run);content.append(para)
        doc._element.body.append(control)
        doc.sections[0].header.paragraphs[0].text='Synthetic header'
        doc.save(path)
        text,report=read_document(str(path))
        self.assertIn('First paragraph.\nSecond paragraph.',text)
        self.assertIn('Before\nAfter',text)
        self.assertEqual(text.count('Controlled text only once'),1)
        self.assertIn('Synthetic header',text)
        self.assertEqual(report['table_count'],1)
    def test_plain_text_disguised_as_word_is_rejected(self):
        path=self.root/'fake.docx';path.write_text(form());self.assert_error(path,'SOURCE_UNSUPPORTED')
    def test_nonword_zip_is_rejected(self):
        path=self.root/'sheet.docx'
        with ZipFile(path,'w') as z: z.writestr('xl/workbook.xml','<workbook/>')
        self.assert_error(path,'SOURCE_UNSUPPORTED')
    def test_truncated_zip_is_rejected(self):
        path=self.root/'broken.docx';path.write_bytes(b'PK\x03\x04broken');self.assert_error(path,'SOURCE_CORRUPT')
    def test_pending_revisions_are_not_merged(self):
        path=self.root/'revisions.docx';doc=Document();p=doc.add_paragraph(SOURCE);p._p.append(OxmlElement('w:del'));doc.save(path);self.assert_error(path,'SOURCE_REVISIONS')
    def test_embedded_content_blocks_incomplete_reading(self):
        path=self.root/'embedded.docx';doc=Document();doc.add_paragraph(SOURCE);doc._element.body.append(OxmlElement('w:altChunk'));doc.save(path);self.assert_error(path,'SOURCE_EMBEDDED_CONTENT')
    def test_empty_word_does_not_pass(self):
        path=self.root/'empty.docx';Document().save(path);self.assert_error(path,'SOURCE_EMPTY')
    def test_native_pdf_reads_all_pages(self):
        path=self.root/'source.pdf'
        with fitz.open() as doc:
            doc.new_page().insert_text((40,60),SOURCE)
            doc.new_page().insert_text((40,60),'Additional source text on page two.')
            doc.save(path)
        text,report=read_document(str(path));self.assertEqual(report['page_count'],2);self.assertIn('page two',text)
    def test_mixed_scanned_pdf_blocks_even_if_first_page_reads(self):
        path=self.root/'mixed.pdf'
        with fitz.open() as doc:
            doc.new_page().insert_text((40,60),SOURCE)
            page=doc.new_page();pix=fitz.Pixmap(fitz.csRGB, (0,0,80,80),False);pix.clear_with(255);page.insert_image(page.rect,pixmap=pix);page.insert_text((40,60),'2');doc.save(path)
        error=self.assert_error(path,'SOURCE_OCR_REQUIRED');self.assertEqual(error.details['pages'],[2])
    def test_encrypted_pdf_requires_unlocked_copy(self):
        path=self.root/'encrypted.pdf'
        with fitz.open() as doc:
            doc.new_page().insert_text((40,60),SOURCE);doc.save(path,encryption=fitz.PDF_ENCRYPT_AES_256,owner_pw='owner',user_pw='reader')
        self.assert_error(path,'SOURCE_ENCRYPTED')
    def test_no_binary_string_salvage_when_legacy_converter_missing(self):
        path=self.root/'legacy.doc';path.write_bytes(bytes.fromhex('D0CF11E0A1B11AE1')+'WordDocument'.encode('utf-16le')+form().encode())
        with patch('utils.doc_parser.shutil.which',return_value=None): self.assert_error(path,'SOURCE_DOC_CONVERSION')
    def test_legacy_conversion_retains_detected_format(self):
        path=self.root/'legacy.doc';path.write_bytes(bytes.fromhex('D0CF11E0A1B11AE1')+'WordDocument'.encode('utf-16le'))
        with patch.object(DocumentParser,'_parse_doc',return_value=form()):
            text,report=read_document(str(path))
        self.assertEqual(report['detected_format'],'doc');self.assertIn(SOURCE,text)

class ExtractionPreflightTests(unittest.TestCase):
    def extract(self,payload):
        from main import extract_document_endpoint
        response=asyncio.run(extract_document_endpoint(type('Request',(),{'json':AsyncMock(return_value=payload)})()))
        return response.status_code,json.loads(response.body)
    def test_unreadable_batch_stops_before_model_and_names_source(self):
        with patch.object(DocumentParser,'parse_with_report',side_effect=SourceError('SOURCE_REVISIONS')),patch('agents.nodes.extraction_node') as model:
            status,data=self.extract({'sources':[{'url':'https://example.invalid/source.docx','name':'draft.docx'}]})
        self.assertEqual(status,422);self.assertEqual(data['source_errors'][0]['code'],'SOURCE_REVISIONS');model.assert_not_called()
    def test_numbered_form_passes_with_literal_facts_and_provenance(self):
        status,data=self.extract({'text':form()})
        self.assertEqual(status,200);self.assertEqual(len(data['matters']),1);self.assertIn(SOURCE,data['matters'][0]['rawNotes']);self.assertEqual(data['ingestion_quality']['status'],'ready_for_review')
    def test_incomplete_matter_never_enters_wizard(self):
        status,data=self.extract({'text':form(summary='')})
        self.assertEqual(status,422);self.assertEqual(data['source_errors'][0]['code'],'SOURCE_INCOMPLETE_MATTERS')
    def test_repeated_labels_within_a_source_are_rejected(self):
        status,data=self.extract({'text':form()+form()})
        self.assertEqual(status,422);self.assertEqual(data['source_errors'][0]['code'],'SOURCE_DUPLICATE_LABELS')
    def test_distinct_repeated_number_blocks_survive_with_original_heading(self):
        source=form()+form().replace(SOURCE,'The team advised on a separate property acquisition.')
        status,data=self.extract({'text':source})
        self.assertEqual(status,200);self.assertEqual(len(data['matters']),2)
        first,second=data['matters']
        self.assertNotEqual(first['id'],second['id'])
        self.assertNotEqual(first['source_label'],second['source_label'])
        self.assertEqual(first['source_heading'],second['source_heading'])
        self.assertNotEqual(first['source_occurrence'],second['source_occurrence'])
        self.assertIn('separate property acquisition',second['source_excerpt'])
        self.assertIn('numbering_reconciliation',data['source_reports'][0])
    def test_same_labels_in_different_sources_do_not_merge_matters(self):
        status,data=self.extract({'sources':[{'name':'first','text':form()},{'name':'second','text':form()}]})
        self.assertEqual(status,200);self.assertEqual(len(data['matters']),2);self.assertNotEqual(data['matters'][0]['id'],data['matters'][1]['id']);self.assertEqual([m['source_document'] for m in data['matters']],['first','second'])
    def test_unstructured_extraction_requires_literal_evidence(self):
        with patch('agents.nodes.extraction_node',return_value={'matters':[{'title':'Invented', 'source_excerpt':'Not in source'}]}):
            status,data=self.extract({'text':SOURCE})
        self.assertEqual(status,422);self.assertEqual(data['source_errors'][0]['code'],'SOURCE_UNGROUNDED_MATTERS')
    def test_mixed_structured_and_unstructured_sources_are_both_extracted(self):
        with patch('agents.nodes.extraction_node',return_value={'matters':[{'title':'Appeal', 'source_excerpt':SOURCE, 'summary':'Model paraphrase'}]}):
            status,data=self.extract({'sources':[{'name':'form','text':form()},{'name':'notes','text':SOURCE}]})
        self.assertEqual(status,200);self.assertEqual(len(data['matters']),2);self.assertEqual(data['matters'][1]['rawNotes'],SOURCE)
    def test_blank_template_slots_are_not_invented_as_matters(self):
        status,data=self.extract({'text':form()+'\nConfidential Matter 2\nE1 Name of client\n\nE2 Summary of matter\n\nE3 Matter value\n'})
        self.assertEqual(status,200);self.assertEqual(len(data['matters']),1);self.assertEqual(data['source_reports'][0]['empty_sections'],['Confidential Matter 2'])
    def test_literal_excerpt_does_not_allow_invented_client_or_amount(self):
        with patch('agents.nodes.extraction_node',return_value={'matters':[{'title':'Appeal', 'client':'Invented Company', 'matter_value':'USD 9000000', 'source_excerpt':SOURCE}]}):
            status,data=self.extract({'text':SOURCE})
        self.assertEqual(status,422);self.assertEqual(data['source_errors'][0]['code'],'SOURCE_UNGROUNDED_MATTERS')
    def test_department_source_is_not_silently_truncated(self):
        narrative='The department advises on tax appeals. ' * 160
        status,data=self.extract({'text':'B10 What is your department best known for\n'+narrative+'\n'+form()})
        self.assertEqual(status,200);self.assertEqual(data['original_b10'],narrative.strip())

class HyperlinkPreflightTests(unittest.TestCase):
    def test_relationship_url_is_preserved_without_fetching_it(self):
        from docx.oxml.ns import qn
        from docx.opc.constants import RELATIONSHIP_TYPE
        with tempfile.TemporaryDirectory() as directory:
            path=Path(directory)/'links.docx';doc=Document();p=doc.add_paragraph('Evidence: ')
            link=OxmlElement('w:hyperlink');link.set(qn('r:id'),p.part.relate_to('https://example.invalid/evidence',RELATIONSHIP_TYPE.HYPERLINK,is_external=True))
            run=OxmlElement('w:r');text=OxmlElement('w:t');text.text='Source article';run.append(text);link.append(run);p._p.append(link);doc.save(path)
            with patch('urllib.request.urlopen') as network:
                text,_=read_document(str(path))
            self.assertIn('Source article (https://example.invalid/evidence)',text);network.assert_not_called()
